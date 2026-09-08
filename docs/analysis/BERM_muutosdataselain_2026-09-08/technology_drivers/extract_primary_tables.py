#!/usr/bin/env python3
"""Read frozen publisher tables into source-labelled records; no network or interpolation.

Extraction dependencies: openpyxl, xlrd. Main offline builder needs only Python stdlib.
The Japanese XLS font size encodes the source's geographic/sample footnote, so it
must be read with formatting_info=True, and those populations must remain separate.
"""
import hashlib
import json
import re
import csv
from pathlib import Path
import openpyxl
import xlrd

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
DATE = "2026-09-08"
out = {"sources": [], "series": []}


def text(fi, en):
    return {"fi": fi, "en": en}


def source(id, title, url, publisher, files, scope, license="See publisher terms; numerical source facts with attribution"):
    artifacts = []
    for name in files:
        p = HERE / name
        b = p.read_bytes()
        artifacts.append({"path": str(p.relative_to(ROOT)), "sha256": hashlib.sha256(b).hexdigest(), "bytes": len(b)})
    out["sources"].append({"id": id, "title": title, "url": url, "publisher": publisher, "retrievedAt": DATE, "license": license, "scope": scope, "artifacts": artifacts})
    return id


def series(id, country, family, title, metric, unit, unit_label, scope, sid, points, limitations, driver=None, status="reported"):
    row = {"id": id, "countryId": country, "familyId": family, "title": title, "metric": metric, "unit": unit, "unitLabel": unit_label, "scope": scope, "sourceIds": [sid], "points": points, "limitations": limitations, "status": status}
    if driver:
        row["driver"] = driver
    out["series"].append(row)


def point(year, value, sid, locator, **extra):
    return {"year": int(year), "value": value, "sourceId": sid, "sourceLocator": locator, "imputed": False, **extra}


def reference(value, label, priority=5):
    return {"normalization": "reference", "referenceValue": value, "referenceLabel": label, "priority": priority, "interpretation": text("Nimetty kiinteä mittakaava ajallisen kuormitusproksin tarkasteluun. Ei paikallinen kenttäamplitudi tai biologinen kynnys.", "Named fixed scale for a temporal system-activity proxy. Not a local field amplitude or biological threshold.")}


def percent(priority=5):
    return {"normalization": "percent", "priority": priority, "interpretation": text("Omistus- tai peittoprosentti jaettuna sadalla. Vastaanottimien omistus ei mittaa lähettimien tehoa, kenttäannosta tai käyttöaikaa.", "Ownership or coverage percentage divided by 100. Receiver ownership does not measure transmitter power, field dose or operating time.")}


# Statistics Finland: actual annual values, preserving all 65 source rows.
sid = source("statfin-electricity-12sv", "StatFin 12sv: Supplies and total consumption of electricity, 1960–2024", "https://pxdata.stat.fi/PxWeb/pxweb/en/StatFin/ehk/12sv.px/", "Statistics Finland", ["statfin_12sv.json", "statfin_12sv_metadata.json", "statfin_12sv_query.json"], text("Suomen sähkön kokonaiskulutus, GWh. Vuosien 1960–1969 lähde on ministeriön Energiatilastot 1960–1975.", "Finland total electricity consumption, GWh. The 1960–1969 records originate in the ministry's Energy Statistics 1960–1975."), "CC BY 4.0; Statistics Finland")
data = json.loads((HERE / "statfin_12sv.json").read_text())
years = sorted(data["dimension"]["timeperiod_y"]["category"]["index"], key=lambda y: data["dimension"]["timeperiod_y"]["category"]["index"][y])
series("fin_electricity_consumption_gwh", "FIN", "electric-grid", text("Sähkön kokonaiskulutus", "Total electricity consumption"), "electricity_consumption_total", "gwh", text("GWh / vuosi", "GWh / year"), out["sources"][-1]["scope"], sid, [point(year, value, sid, f"Table 12sv; timeperiod_y={year}; energia_30_20200915=SSS; contentscode=maara_gwh") for year, value in zip(years, data["value"]) if value is not None], [text("Sisältää teollisuuden, palvelut, kotitaloudet ja verkon häviöt. Ei kotitalouden keskimääräinen kulutus tai paikallinen sähkö- tai magneettikenttä.", "Includes industry, services, households and network losses. Not mean household consumption or a local electric or magnetic field."), text("Tilastokeskus lisää pien-CHP-tuotannon; tästä syystä kokonaismäärä eroaa Energiateollisuuden sarjasta.", "Statistics Finland includes micro-CHP production; totals therefore differ from Finnish Energy Industries figures.")], reference(100000, text("100 000 GWh / vuosi", "100,000 GWh / year"), 1))

# US Census compilation. Only numeric source cells; NA is never a zero.
sid = source("census-hs42-media", "U.S. Census Statistical Abstract 2003, Table HS-42: Selected Communications Media 1920–2001", "https://www2.census.gov/library/publications/2004/compendia/statab/123ed/hist/hs-42.pdf", "U.S. Census Bureau, compiled from RAB, TVB and FCC sources", ["census_hs42.pdf", "census_hs42.txt"], text("Radio Advertising Bureaun ja Television Bureau of Advertisingin kotitalousarvioiden sekä TV-asematilastojen virallinen historiakooste.", "Official historical compilation of Radio Advertising Bureau and Television Bureau of Advertising household estimates and TV station statistics."))
rows = [line.split() for line in (HERE / "census_hs42.txt").read_text().splitlines() if re.match(r"^\d{4}\s+\.{2,}", line)]
for index, key, family, name in [(3, "radio", "radio-broadcast", text("Radio kotitaloudessa", "Households with radio")), (4, "tv", "terrestrial-tv", text("Televisio kotitaloudessa", "Households with television"))]:
    points = [point(r[0], float(r[index]), sid, f"HS-42, year label {r[0]}, {'Radio' if key == 'radio' else 'Television'} column; footnote {2 if key == 'radio' else 3}") for r in rows if r[index] != "(NA)"]
    scope = text("Yhdysvaltain kotitaloudet; vuoden lopun arviot.", "U.S. households; year-end estimates.") if key == "radio" else text("Yhdysvaltain kotitaloudet ilman Alaskaa ja Havaijia. TV-vuosiluvut 1970–1979 viittaavat edellisen vuoden syyskuuhun; muut tammikuuhun.", "U.S. households excluding Alaska and Hawaii. TV year labels 1970–1979 refer to September of the preceding year; others to January.")
    series(f"usa_{key}_household_ownership", "USA", family, name, f"{key}_household_ownership", "percent", text("% kotitalouksista", "% of households"), scope, sid, points, [text("Omistustilasto ei erottele käyttöaikaa, vastaanottimen rakennetta, antenniverkon peittoa eikä lähetystehoa. Vastaanotin ei ole sama asia kuin RF-lähetin.", "Ownership does not identify operating time, receiver design, terrestrial coverage or transmitter power. A receiver is not an RF transmitter.")], percent(10), "estimate")
station_points = [point(r[0], sum(0 if v == "-" else int(v.replace(",", "")) for v in r[7:9]), sid, f"HS-42, year {r[0]}, sum of VHF and UHF commercial television station columns; footnote 5") for r in rows if "(NA)" not in r[7:9]]
series("usa_commercial_tv_stations", "USA", "terrestrial-tv", text("Kaupalliset TV-asemat (VHF + UHF)", "Commercial TV stations (VHF + UHF)"), "commercial_tv_stations", "stations", text("asemaa", "stations"), text("Yhdysvaltain kaupalliset VHF- ja UHF-televisioasemat, HS-42:n sarakkeiden summa. Ei kaikkia lähettimiä tai toistimia.", "U.S. commercial VHF and UHF television stations, sum of HS-42 columns. Not all transmitters or repeaters."), sid, station_points, [text("Asemien määrä ei sisällä tehoa, antennia, taajuuskaistaa, käyttöaikaa tai asutuksen etäisyyttä. Vuodesta 1997 lähteenä FCC; viitepäivä vaihtuu.", "Station counts contain no power, antenna, band, operating time or residential distance. FCC is the source from 1997; reference dates change.")], reference(1000, text("1 000 kaupallista TV-asemaa", "1,000 commercial TV stations"), 1))

# UK: do not splice former geographical or sector definitions into a single series.
sid = source("desnz-historical-electricity", "Historical electricity data: 1920–2025, July 2026 release", "https://www.gov.uk/government/statistical-data-sets/historical-electricity-data", "UK Department for Energy Security and Net Zero", ["gb_electricity_1920.xlsx"], text("DUKESin historialliset sähkötilastot. Tässä tuodaan todelliset taulukkosolut vuoteen 2024 asti; erillistä arvioidun tuotannon välilehteä ei käytetä.", "DUKES historical electricity statistics. Actual table cells are imported through 2024; the separate estimated-generation worksheet is not used."), "Open Government Licence v3.0")
w = openpyxl.load_workbook(HERE / "gb_electricity_1920.xlsx", data_only=True)
s = w["Generated and Supplied"]
for start, end, suffix, geography in [(1920, 1950, "gb_1920_1950", text("Iso-Britannia ilman Pohjois-Irlantia", "Great Britain excluding Northern Ireland")), (1951, 2024, "uk_1951_2024", text("Yhdistynyt kuningaskunta", "United Kingdom"))]:
    points = [point(s[f"A{i}"].value, s[f"B{i}"].value, sid, f"Generated and Supplied!B{i}; year=A{i}; Major power producers; notes 1,6") for i in range(8, s.max_row + 1) if isinstance(s[f"A{i}"].value, int) and start <= s[f"A{i}"].value <= end and isinstance(s[f"B{i}"].value, (int, float))]
    series(f"gbr_mpp_generation_{suffix}", "GBR", "electric-grid", text(f"Suurten sähköntuottajien tuotanto ({start}–{end})", f"Major power producers' generation ({start}–{end})"), "mpp_electricity_generation", "gwh", text("GWh / vuosi", "GWh / year"), geography, sid, points, [text("Suurten sähköntuottajien bruttotuotanto. Ei koko maan kulutus. Tuuli-MPP:t sisältyvät vuodesta 2007 ja aurinko-MPP:t vuodesta 2015; muut tuottajat ovat erillinen luokka.", "Gross generation by major power producers. Not total national consumption. Wind MPPs enter from 2007 and solar MPPs from 2015; other generators form a separate category."), text("Ennen vuotta 1951 alue on Iso-Britannia. Aluerajan vuoksi historialliset jaksot säilytetään erillisinä sarjoina.", "Before 1951 the territory is Great Britain. The historical periods remain separate series because the territory changes.")], reference(100000, text("100 000 GWh / vuosi", "100,000 GWh / year"), 3))
s = w["Supply, Availability & Consumpt"]
points = [point(s[f"A{i}"].value, s[f"J{i}"].value * 1000, sid, f"Supply, Availability & Consumpt!J{i}; year=A{i}; Domestic consumption; TWh × 1000 = GWh") for i in range(7, s.max_row + 1) if isinstance(s[f"A{i}"].value, int) and s[f"A{i}"].value <= 2024 and isinstance(s[f"J{i}"].value, (int, float))]
# Year labels containing a source footnote remain genuine observations. A second
# annotated 1987 row is a historical comparison definition: prefer the unannotated
# row already selected above, and expose this scope choice in metadata.
for i in range(7, s.max_row + 1):
    match = re.fullmatch(r"(\d{4}) \[note 7\]", str(s[f"A{i}"].value))
    if match and not any(p["year"] == int(match.group(1)) for p in points):
        points.append(point(int(match.group(1)), s[f"J{i}"].value * 1000, sid, f"Supply, Availability & Consumpt!J{i}; A{i}={s[f'A{i}'].value}; Domestic consumption; TWh × 1000 = GWh"))
points.sort(key=lambda p: p["year"])
series("gbr_domestic_electricity_consumption", "GBR", "electric-grid", text("Kotitalouksien sähkönkulutus", "Domestic electricity consumption"), "domestic_electricity_consumption", "gwh", text("GWh / vuosi", "GWh / year"), text("DUKESin Yhdistyneen kuningaskunnan kotitaloussektori. Vuosiluvut vain taulukon havaintovuosille.", "DUKES United Kingdom domestic sector. Years only where the table reports a value."), sid, points, [text("Kotitaloussektorin kokonaiskulutus, ei asukaskohtainen arvo. Varhaisessa osassa on vuosiaukkoja; niitä ei täytetä.", "Total domestic-sector consumption, not a per-person value. Early missing years remain missing."), text("Huomautus7 käsittelee tuottajajoukon laajenemista. Vuoden1986 huomautettu rivi säilytetään. Vuodelle1987 työkirjassa on kaksi riviä: tässä käytetään huomautuksetonta93,25TWh-riviä; huomautettua91,83TWh-vertailuriviä ei yhdistetä toiseksi saman vuoden havainnoksi.", "Note7 concerns expanded producer coverage. The annotated 1986 row is retained. The workbook has two 1987 rows: this series uses the unannotated 93.25TWh row; the annotated 91.83TWh comparison row is not added as a second observation of the same year.")], reference(100000, text("100 000 GWh / vuosi", "100,000 GWh / year"), 1))

# GB meter stock in smart mode: Q4 year-end points only, large suppliers throughout.
sid = source("desnz-smart-meters-2024", "Smart meters in Great Britain, quarterly update March 2024: data tables", "https://www.gov.uk/government/statistics/smart-meters-in-great-britain-quarterly-update-march-2024", "UK Department for Energy Security and Net Zero", ["gb_smart_2024.xlsx"], text("Ison-Britannian sähkö- ja kaasumittarit. Tässä vain suurten toimittajien kotitalouksien sähkömittarit Q4-vuoden lopussa.", "Great Britain electricity and gas meters. This extraction uses domestic electricity meters operated by large suppliers at Q4 year-end only."), "Open Government Licence v3.0")
w = openpyxl.load_workbook(HERE / "gb_smart_2024.xlsx", data_only=True)
s = w["Table1"]
points = []
for i in range(9, s.max_row + 1):
    match = re.fullmatch(r"Q4 (\d{4})", str(s[f"A{i}"].value))
    if match:
        year = int(match.group(1))
        points.append(point(year, s[f"F{i}"].value, sid, f"Table1!F{i}; A{i}=Q4 {year}; electricity smart in smart mode; denominator H{i}+I{i}", asOf=f"{year}-12-31", denominator={"value": s[f"H{i}"].value + s[f"I{i}"].value, "unit": "meters", "population": text("Saman taulukkorivin suurten toimittajien kotitalouksien kaikki sähkömittarit", "All large-supplier domestic electricity meters in the same table row")}))
series("gbr_domestic_electricity_smart_mode", "GBR", "smart-metering", text("Älytilassa toimivat kotitalouksien sähkömittarit", "Domestic electricity meters operating in smart mode"), "electricity_smart_mode_stock", "meters", text("sähkömittaria", "electricity meters"), out["sources"][-1]["scope"], sid, points, [text("Suurten toimittajien joukko muuttuu lähteen huomautusten mukaisesti. Perinteisessä tilassa toimivat älymittarit eivät kuulu osoittajaan. Ei kumulatiivinen asennusmäärä.", "The large-supplier group changes as documented in source notes. Smart meters in traditional mode are excluded from the numerator. Not cumulative installations."), text("Älytila ei määritä protokollaa, lähetystehoa tai yksilön altistusta.", "Smart mode does not identify protocol, transmitting power or individual exposure.")], {"normalization": "ratio", "priority": 1, "interpretation": text("Älytilassa toimivien mittarien osuus saman rajauksen kaikista sähkömittareista.", "Share operating in smart mode among all electricity meters in the same scope.")})

# Japan: the original small-font cells explicitly mean non-farm households.
sid = source("esri-durables-2004", "Consumer Confidence Survey: historical diffusion of principal consumer durables, March 2004", "https://www.esri.cao.go.jp/jp/stat/shouhi/0403fukyuritsu.xls", "Economic and Social Research Institute, Cabinet Office, Japan", ["jp_durables_2004.xls"], text("Alkuperäinen historiallinen kotitalouksien laiteomistustaulukko. Pieni kirjasin tarkoittaa ei-maatalouskotitalouksia ja vuoteen 1963 asti vain kaupunkeja; aineistot erotetaan.", "Original historical household durable-ownership table. Small font identifies non-farm households, and through 1963 urban areas only; populations are separated."))
w = xlrd.open_workbook(str(HERE / "jp_durables_2004.xls"), formatting_info=True)
s = w.sheet_by_index(0)
for col, key, label, family in [(33, "monochrome_tv", text("Mustavalkotelevisio kotitaloudessa", "Households with monochrome television"), "terrestrial-tv"), (30, "colour_tv", text("Väritelevisio kotitaloudessa", "Households with colour television"), "terrestrial-tv"), (26, "air_conditioner", text("Ilmastointilaite kotitaloudessa", "Households with room air conditioning"), "power-conversion")]:
    groups = {}
    era = 1925
    for i in range(4, 52):
        if s.cell_value(i, 0) == "平成":
            era = 1988
        value = s.cell_value(i, col)
        if not isinstance(value, (int, float)):
            continue
        raw_date = str(s.cell_value(i, 1))
        year = era + int(raw_date.split(".")[0])
        month = int(raw_date.split(".")[1])
        small = w.font_list[w.xf_list[s.cell_xf_index(i, col)].font_index].height < 180
        group = "urban_nonfarm" if small and year <= 1963 else "nonfarm" if small else "general"
        groups.setdefault(group, []).append(point(year, value, sid, f"Worksheet 1, {xlrd.formula.colname(col)}{i+1}; Japanese era date {raw_date}; font height={w.font_list[w.xf_list[s.cell_xf_index(i, col)].font_index].height}; footnotes rows53–54", asOf=f"{year}-{month:02d}"))
    for group, points in groups.items():
        pop = {"urban_nonfarm": text("Kaupunkien ei-maatalouskotitaloudet, Japani", "Urban non-farm households, Japan"), "nonfarm": text("Ei-maatalouskotitaloudet, Japani", "Non-farm households, Japan"), "general": text("Japanin yleiset kotitaloudet (vastaa nykyisen tutkimuksen vähintään kahden hengen kotitalouksia)", "Japanese general households (corresponding to the modern survey's households of two or more persons)")}[group]
        limitations = [text("Omistusosuus ei osoita laitteen päivittäistä käyttöä, sähköistä rakennetta tai kenttää. Mustavalko- ja väritelevision osuuksia ei lasketa yhteen: samalla kotitaloudella voi olla molemmat.", "Ownership share does not establish daily use, electrical design or fields. Monochrome and colour TV shares are not added: a household may own both."), text("Havaintokuukausi säilyy asOf-kentässä. Eri kotitalousrajaukset muodostavat erilliset sarjat.", "Observation month is retained in asOf. Different household populations form separate series.")]
        if key == "air_conditioner":
            limitations.append(text("Kaikki huoneilmastointilaitteet; ei invertterilaitteiden osuus. Sarjaa ei oletusarvoisesti muunneta power-conversion-perheen ajuriksi.", "All room air conditioners; not the share using inverters. This series is not assigned a default power-conversion driver."))
        series(f"jpn_{key}_{group}", "JPN", family, {lang: label[lang] + " — " + pop[lang] for lang in ["fi", "en"]}, f"{key}_household_ownership", "percent", text("% rajatuista kotitalouksista", "% of scoped households"), pop, sid, points, limitations, percent(5) if family == "terrestrial-tv" else None, "estimate")

# US residential sales offer a long measured load history without splicing the
# pre-1989 utility-only generation definition into modern all-sector generation.
sid = source("eia-residential-electricity", "Monthly Energy Review August 2026, Table 7.6: Residential electricity sales", "https://www.eia.gov/totalenergy/data/monthly/#electricity", "U.S. Energy Information Administration", ["eia_end_use.csv", "eia_electricity_notes.pdf"], text("Sähkön myynti asuinsektorin loppuasiakkaille 50 osavaltiossa ja District of Columbiassa. Palveluluokka, ei kotitalouksien määrä.", "Electricity sales to residential ultimate customers in the 50 states and District of Columbia. A class of service, not the number of households."), "U.S. government public data; attribution to EIA")
rows = list(csv.DictReader((HERE / "eia_end_use.csv").open()))
points = [point(int(row["YYYYMM"][:4]), float(row["Value"]), sid, f"MER Table7.6 CSV; MSN=ESRCPUS; YYYYMM={row['YYYYMM']} (13=annual total); million kWh = GWh") for row in rows if row["MSN"] == "ESRCPUS" and row["YYYYMM"].endswith("13") and int(row["YYYYMM"][:4]) <= 2024]
series("usa_residential_electricity_sales", "USA", "electric-grid", text("Asuinsektorille myyty sähkö", "Residential electricity sales"), "residential_electricity_sales", "gwh", text("GWh / vuosi", "GWh / year"), out["sources"][-1]["scope"], sid, points, [text("Sähköyhtiöiden raportoimat loppuasiakasmyynnit; vuodesta 1996 myös muiden energiapalveluntarjoajien myynti. Oma tuotanto ei ole sähkömyyntiä.", "Utility-reported ultimate-customer sales; other energy service providers included from 1996. Self-generation is not electricity sales."), text("Luokitus seuraa palveluluokkaa ja sisältää siinä laskutetun sähköautojen latauksen. GWh kuvaa vuosittaista energiavirtaa, ei paikallista kenttää tai laitekannan kokoa.", "Classification follows class of service and includes electric-vehicle charging billed in that class. GWh measures annual energy flow, not a local field or equipment stock.")], reference(100000, text("100 000 GWh / vuosi", "100,000 GWh / year"), 1))

# Federal Republic of Germany, original EVS publication (not present-day territory).
# Verified against printed page 480, PDF page 26, Table 4. Its three reference
# years and integer percentages are transcribed exactly; no annual values filled.
sid = source("destatis-evs-1974", "Wirtschaft und Statistik 7/1974, Table 4: Ausstattung privater Haushalte 1962, 1969 und 1973", "https://www.destatis.de/GPStatistik/servlets/MCRFileNodeServlet/DEAusgabe_derivate_00000708/Wirtschaft_und_Statistik-1974-07.pdf", "Statistisches Bundesamt", ["de_wista_1974_07.pdf", "de_wista_1974_07.txt"], text("Saksan liittotasavallan historiallinen EVS-otos, ei nykyisen yhdistyneen Saksan alue. Ulkomaalais- ja laitoskotitaloudet rajattu pois; 1969 ja 1973 myös ylimmät tuloluokat.", "Historical Federal Republic of Germany EVS sample, not the territory of reunified Germany. Foreign-national and institutional households excluded; top income groups also excluded in 1969 and 1973."))
for key, family, label, values in [("tv", "terrestrial-tv", text("Televisio kotitaloudessa, Länsi-Saksa", "Households with television, West Germany"), [34, 73, 87]), ("radio", "radio-broadcast", text("Radio kotitaloudessa, Länsi-Saksa", "Households with radio, West Germany"), [79, 83, 86])]:
    series(f"deu_west_{key}_ownership", "DEU", family, label, f"{key}_household_ownership", "percent", text("% rajatuista kotitalouksista", "% of scoped households"), out["sources"][-1]["scope"], sid, [point(year, value, sid, f"Printed p480, PDF p26, Table4, {'Fernsehgerät' if key == 'tv' else 'Rundfunkgerät'} row, {year} column") for year, value in zip([1962, 1969, 1973], values)], [text("1969 rajataan pois vähintään 10 000 DM:n ja 1973 vähintään 15 000 DM:n kuukausitulon kotitaloudet. Televisiot eivät sisällä phonokombinaatioiden laitteita.", "Households with monthly net income at least DM10,000 in 1969 and DM15,000 in 1973 are excluded. Television sets in phono combinations are excluded."), text("Omistusosuus ei mittaa RF-lähetyksiä tai vastaanottimien käyttöaikaa. Maavalinta DEU näyttää tämän historiallisen alarajauksen, ei harmonisoi sitä nykyiseen väestöön.", "Ownership share does not measure RF transmissions or receiver operating time. The DEU country selection displays this historical subset without harmonising it to the modern population.")], percent(5), "estimate")

dest = HERE / "primary_extract.json"
dest.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n")
print(f"Extracted {len(out['series'])} series, {sum(len(s['points']) for s in out['series'])} points")
